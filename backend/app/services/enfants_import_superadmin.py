"""Import CSV ou Excel (.xlsx) des enfants « codifiés » par le super admin — aligné sur le flux existant (`Enfant` + `auto_sync_enfants_eligibles_du_parent`)."""

from __future__ import annotations

import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any

from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.models.enums import LienParente, Sexe, UserRole
from app.models.models import Enfant, User
from app.services.inscriptions import auto_sync_enfants_eligibles_du_parent, _date_naissance_dans_plage

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "matricule_parent": ("matricule_parent", "matricule", "matricule_agent"),
    "prenom": ("prenom", "prenom_enfant"),
    "nom": ("nom", "nom_enfant"),
    "date_naissance": ("date_naissance", "naissance", "date_de_naissance"),
    "sexe": ("sexe",),
    "lien_parente": ("lien_parente", "lien", "lien_de_parente"),
}


def _slug_header(h: str) -> str:
    return re.sub(r"\s+", "_", (h or "").strip().lower().lstrip("\ufeff"))


def _detect_delimiter(sample: str) -> str:
    first = sample.split("\n", 1)[0] if sample else ""
    if first.count(";") > first.count(","):
        return ";"
    return ","


def _parse_date_cell(val: str) -> date:
    v = (val or "").strip()
    if not v:
        raise ValueError("Date de naissance vide.")
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Date non reconnue : {val!r} (formats acceptés : AAAA-MM-JJ ou JJ/MM/AAAA).")


def _parse_sexe(val: str) -> Sexe:
    v = (val or "").strip().upper()
    if v in ("M", "MASCULIN", "GARCON", "GARÇON"):
        return Sexe.M
    if v in ("F", "FEMININ", "FÉMININ", "FILLE"):
        return Sexe.F
    raise ValueError(f"Sexe invalide : {val!r} (M ou F).")


def _parse_lien(val: str) -> LienParente:
    raw = (val or "").strip().upper().replace(" ", "_")
    if raw in ("PERE", "MERE", "TUTEUR_LEGAL", "AUTRE"):
        return LienParente(raw)
    v = re.sub(r"[éèê]", "e", (val or "").strip().lower())
    v = " ".join(v.split())
    mapping = {
        "pere": LienParente.PERE,
        "mere": LienParente.MERE,
        "tuteur legal": LienParente.TUTEUR_LEGAL,
        "tuteur": LienParente.TUTEUR_LEGAL,
        "autre": LienParente.AUTRE,
    }
    if v in mapping:
        return mapping[v]
    raise ValueError(f"Lien de parenté invalide : {val!r} (PERE, MERE, TUTEUR_LEGAL, AUTRE ou libellés usuels).")


def _norm_name(s: str) -> str:
    return " ".join((s or "").split()).lower()


def _resolve_columns(headers: list[Any]) -> dict[str, int] | None:
    slugs = [_slug_header(str(h) if h is not None else "") for h in headers]
    out: dict[str, int] = {}
    for key, aliases in _HEADER_ALIASES.items():
        idx = None
        for a in aliases:
            try:
                idx = slugs.index(a)
                break
            except ValueError:
                continue
        if idx is None:
            return None
        out[key] = idx
    return out


def _trim_trailing_blanks_row(row: list[Any]) -> list[Any]:
    cells = list(row)
    while cells and (cells[-1] is None or cells[-1] == ""):
        cells.pop()
    return cells


def _xlsx_bytes_to_rows(raw: bytes) -> list[list[Any]]:
    """Retourne les lignes avec types openpyxl intacts (int, float, datetime…).

    Ne pas convertir les entiers en datesici : un matricule parent saisi comme
    nombre (ex. 7777) est un serial Excel « date » faux positif (ex. 1921-04-16).
    La colonne date_naissance est interprétée dans `_coerce_cell_for_import`.
    """
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb.active
        raw_rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            cells = _trim_trailing_blanks_row(list(row))
            if any(c is not None and c != "" for c in cells):
                raw_rows.append(cells)
        return raw_rows
    finally:
        wb.close()


def _coerce_cell_for_import(val: Any, field: str) -> str:
    """Normalise une cellule CSV/Excel selon le rôle de la colonne."""
    if val is None:
        return ""
    if field == "matricule_parent":
        if isinstance(val, float) and not isinstance(val, bool):
            if val.is_integer():
                return str(int(val))
            return str(val).strip()
        if isinstance(val, int) and not isinstance(val, bool):
            return str(val)
        return str(val).strip()
    if field == "date_naissance":
        if isinstance(val, datetime):
            return val.date().strftime("%Y-%m-%d")
        if isinstance(val, date):
            return val.strftime("%Y-%m-%d")
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            try:
                from openpyxl.utils.datetime import from_excel

                return from_excel(val).date().strftime("%Y-%m-%d")
            except Exception:
                s = str(val).strip()
                if s.endswith(".0"):
                    s = s[:-2]
                return s
        return str(val).strip()
    if isinstance(val, datetime):
        return val.date().strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, float) and not isinstance(val, bool) and val.is_integer():
        return str(int(val))
    if isinstance(val, int) and not isinstance(val, bool):
        return str(val)
    return str(val).strip()


def _process_enfants_import_rows(db: Session, raw_rows: list[list[Any]]) -> dict[str, Any]:
    if len(raw_rows) < 2:
        return {
            "ok": False,
            "error": "Le fichier doit contenir une ligne d'en-têtes et au moins une ligne de données.",
        }

    col = _resolve_columns(raw_rows[0])
    if col is None:
        expected = ", ".join(sorted({a for al in _HEADER_ALIASES.values() for a in al}))
        return {
            "ok": False,
            "error": (
                f"En-têtes invalides. Colonnes reconnues (au choix) : {expected}. "
                "CSV : séparateur virgule ou point-virgule. Excel : première feuille, ligne 1 = en-têtes."
            ),
        }

    results: list[dict[str, Any]] = []
    parents_to_sync: dict[int, User] = {}

    for ligne_no, parts in enumerate(raw_rows[1:], start=2):
        def cell(slug: str) -> str:
            i = col[slug]
            raw = parts[i] if i < len(parts) else None
            return _coerce_cell_for_import(raw, slug) or ""

        mat = cell("matricule_parent")
        prenom = cell("prenom")
        nom = cell("nom")
        date_s = cell("date_naissance")
        sexe_s = cell("sexe")
        lien_s = cell("lien_parente")

        if not any([mat, prenom, nom, date_s, sexe_s, lien_s]):
            continue

        if not mat:
            results.append({"ligne": ligne_no, "ok": False, "message": "Matricule parent manquant."})
            continue
        if not prenom or not nom:
            results.append({"ligne": ligne_no, "ok": False, "message": "Prénom ou nom enfant manquant."})
            continue

        try:
            dna = _parse_date_cell(date_s)
        except ValueError as e:
            results.append({"ligne": ligne_no, "ok": False, "message": str(e)})
            continue

        if not _date_naissance_dans_plage(dna):
            results.append(
                {"ligne": ligne_no, "ok": False, "message": "Année de naissance hors plage 2012–2019."},
            )
            continue

        try:
            sexe = _parse_sexe(sexe_s)
        except ValueError as e:
            results.append({"ligne": ligne_no, "ok": False, "message": str(e)})
            continue

        try:
            lien = _parse_lien(lien_s)
        except ValueError as e:
            results.append({"ligne": ligne_no, "ok": False, "message": str(e)})
            continue

        u = (
            db.query(User)
            .options(joinedload(User.parent_profile))
            .filter(func.lower(User.matricule) == mat.strip().lower(), User.role == UserRole.PARENT)
            .first()
        )
        if u is None or u.parent_profile is None:
            results.append(
                {
                    "ligne": ligne_no,
                    "ok": False,
                    "message": f'Aucun compte parent trouvé pour le matricule « {mat.strip()} ».',
                },
            )
            continue

        parent = u.parent_profile
        existing = db.query(Enfant).filter(Enfant.parent_id == parent.id).all()

        np, nn = _norm_name(prenom), _norm_name(nom)
        if any(
            _norm_name(e.prenom) == np and _norm_name(e.nom) == nn and e.date_naissance == dna for e in existing
        ):
            results.append(
                {
                    "ligne": ligne_no,
                    "ok": False,
                    "message": "Doublon : cet enfant existe déjà pour ce parent.",
                },
            )
            continue

        is_first = len(existing) == 0
        enfant = Enfant(
            parent_id=parent.id,
            prenom=prenom.strip()[:191],
            nom=nom.strip()[:191],
            date_naissance=dna,
            sexe=sexe,
            lien_parente=lien,
            is_titulaire=is_first,
        )
        db.add(enfant)
        db.flush()
        parents_to_sync[int(u.id)] = u
        results.append(
            {
                "ligne": ligne_no,
                "ok": True,
                "message": f'Enfant enregistré : {prenom.strip()} {nom.strip()}',
                "matricule": mat.strip(),
            },
        )

    for user_obj in parents_to_sync.values():
        auto_sync_enfants_eligibles_du_parent(db=db, user=user_obj)

    if not results:
        return {"ok": False, "error": "Aucune ligne de données exploitable (lignes vides ?)."}

    created = sum(1 for r in results if r.get("ok"))
    errors = [r for r in results if not r.get("ok")]
    return {
        "ok": True,
        "lignes_traitees": len(results),
        "creees": created,
        "erreurs": len(errors),
        "details": results[:300],
    }


def import_enfants_csv_superadmin(db: Session, csv_text: str) -> dict[str, Any]:
    """Import depuis texte CSV (UTF-8)."""
    text = (csv_text or "").lstrip("\ufeff").strip()
    if not text:
        return {"ok": False, "error": "Fichier vide."}

    delim = _detect_delimiter(text[:4096])
    raw_rows = list(csv.reader(StringIO(text), delimiter=delim))
    return _process_enfants_import_rows(db, raw_rows)


def import_enfants_xlsx_superadmin(db: Session, raw: bytes) -> dict[str, Any]:
    """Import depuis classeur Excel (.xlsx) : première feuille active, ligne 1 = en-têtes."""
    try:
        raw_rows = _xlsx_bytes_to_rows(raw)
    except Exception as e:
        return {
            "ok": False,
            "error": f"Lecture du fichier Excel impossible : {e}. Utilisez le format .xlsx (Excel 2007+).",
        }
    return _process_enfants_import_rows(db, raw_rows)


def import_enfants_fichier_superadmin(db: Session, raw: bytes, filename: str | None = None) -> dict[str, Any]:
    """
    Dispatch : .xlsx / .xlsm → openpyxl ; sinon traité comme CSV UTF-8.
    L’ancien format .xls n’est pas pris en charge.
    """
    fn = (filename or "").lower()
    if fn.endswith((".xlsx", ".xlsm")):
        return import_enfants_xlsx_superadmin(db, raw)
    if fn.endswith(".xls"):
        return {
            "ok": False,
            "error": "Le format .xls n’est pas pris en charge. Enregistrez le classeur au format .xlsx dans Excel.",
        }
    # Sans extension fiable : ZIP « Office » = xlsx
    if len(raw) >= 4 and raw[:4] == b"PK\x03\x04" and not fn.endswith(".csv"):
        return import_enfants_xlsx_superadmin(db, raw)
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return {
            "ok": False,
            "error": "Fichier non reconnu : utilisez un CSV UTF-8 ou un classeur Excel .xlsx.",
        }
    return import_enfants_csv_superadmin(db, text)
